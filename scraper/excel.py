"""Mesečni Excel u obliku istorijske tabele: jedan sheet po seriji (15a, 15b, … 51b), red po satu.

    python -m scraper.excel --mesec 2026-10   # -> podaci/excel/bicikli_2026-10.xlsx

Kolone sheeta: Date | Time (01:00–24:00) | y_A | razlog_A | oznaka
  (stari dvosmerni brojači 25a, 26, 28, 40a/b, 41, 46a/b, 47, 51a/b: Date | Time | A01 | A02 | Sum | razlog_A | oznaka,
   A01 = smer 1x, A02 = smer 2x, Sum = A01 + A02 = y_A, kao u istorijskom masteru)
  - y_A = broj biciklista u satu (cilj modela; feed = A, samo bicikli);
  - razlog_A = „ok“ za ispravan sat, inače razlog (nema_citanja, zastarelo, nema_podataka, …), y_A prazno;
  - oznaka = neobičan sat po smeru (anomalije.py), vrednost se ne menja.
Ćelije sa vrednošću > 400 su crvene (kao u istorijskoj tabeli); neobični sati (anomalije.py) su
narandžasti (nule koje liče na kvar senzora) ili žuti (nagli izolovani skok); nedostajući sati su sivi.
Klase vozila (A/B/C), OCC i GAP iz istorijskog izvoza sajt ne objavljuje, pa ih nema.
"""
from __future__ import annotations

import argparse
import logging
import sys
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import anomalije, konfig
from .parsiranje import OCEKIVANE_SERIJE

log = logging.getLogger(__name__)

PRAG_CRVENO = 400
_ZAGLAVLJE_FILL = PatternFill("solid", fgColor="305496")
_ZAGLAVLJE_FONT = Font(bold=True, color="FFFFFF")
_CRVENO_FILL = PatternFill("solid", fgColor="FF0000")
_CRVENO_FONT = Font(bold=True, color="FFFFFF")
_OZNAKA_FILL = {
    "nula_dan": PatternFill("solid", fgColor="FFC000"),   # narandžasto
    "nula_niz": PatternFill("solid", fgColor="FFC000"),
    "skok": PatternFill("solid", fgColor="FFFF00"),       # žuto
}
_NEMA_FILL = PatternFill("solid", fgColor="D9D9D9")        # sivo — nedostaje vrednost


def folder_excel(koren: Path | None = None) -> Path:
    return (koren or konfig.PODACI) / "excel"


def _vreme(ts: pd.Timestamp) -> str:
    """Kraj sata kao tekst 01:00–24:00 (00:00 sledećeg dana -> 24:00)."""
    return f"{ts.hour or 24:02d}:00"


_LOKACIJA_FILL = PatternFill("solid", fgColor="C00000")   # tamnocrveno — opis lokacije, nije podatak o biciklistima
_LOKACIJA_FONT = Font(bold=True, color="FFFFFF")


def opis_stanica(koren: Path | None = None) -> pd.DataFrame:
    """Iz poslednjeg čitanja: serija, broj mesta, naziv (ulica), smerovi sa opisom, koordinate."""
    fajlovi = sorted((koren or konfig.PODACI).glob("[0-9][0-9][0-9][0-9]/[0-9][0-9]/[0-9][0-9]/[0-9]*.csv"))
    if not fajlovi:
        return pd.DataFrame(columns=["serija", "mesto", "naziv", "smerovi", "lat", "lon"])
    c = pd.read_csv(fajlovi[-1], dtype=str)
    redovi = []
    for serija, g in c.groupby("serija"):
        g = g.sort_values("direction")
        naziv = g["locationID"].iloc[0].split(" ", 1)[1] if " " in g["locationID"].iloc[0] else g["locationID"].iloc[0]
        redovi.append({"serija": serija, "mesto": g["location"].iloc[0], "naziv": naziv.strip(),
                       "smerovi": "; ".join(f"smer {d}: {o}" for d, o in zip(g["direction"], g["directionDesc"])),
                       "lat": float(g["lat"].iloc[0]), "lon": float(g["lon"].iloc[0])})
    redosled = {s: i for i, s in enumerate(OCEKIVANE_SERIJE)}
    return pd.DataFrame(redovi).sort_values("serija", key=lambda x: x.map(redosled).fillna(99)).reset_index(drop=True)


def _ucitaj(folder_meseca: Path, ime: str) -> pd.DataFrame:
    fajlovi = sorted(folder_meseca.glob(f"[0-9][0-9]/{ime}"))
    if not fajlovi:
        return pd.DataFrame()
    df = pd.concat([pd.read_csv(f, dtype={"direction": str, "stanica_id": str, "serija": str}) for f in fajlovi],
                   ignore_index=True)
    df["ts"] = pd.to_datetime(df["ts"])
    df["ts_utc"] = pd.to_datetime(df["ts_utc"], utc=True)
    return df


def napravi_mesec(d: date, koren: Path | None = None) -> Path | None:
    folder = konfig.folder_meseca(d, koren)
    serije = _ucitaj(folder, "satno.csv")
    smerovi = _ucitaj(folder, "satno_smerovi.csv")
    if serije.empty:
        return None

    opisi = opis_stanica(koren).set_index("serija")
    wb = Workbook()
    wb.remove(wb.active)
    redosled = [s for s in OCEKIVANE_SERIJE if s in set(serije["stanica_id"])]
    redosled += sorted(set(serije["stanica_id"]) - set(redosled))
    for serija in redosled:
        sr = serije[serije["stanica_id"] == serija].sort_values("ts_utc")
        sm = smerovi[smerovi["serija"] == serija].reset_index(drop=True)
        smerovi_serije = sorted(sm["direction"].unique())
        dvosmerna = len(smerovi_serije) > 1
        po_smeru = {s: g.set_index("ts_utc")["y_A"] for s, g in sm.groupby("direction")}
        oznake_smera = {}
        for s, g in sm.groupby("direction"):
            o = anomalije.oznaci(g)
            oznake_smera[s] = pd.Series(o.values, index=g.loc[o.index, "ts_utc"])

        # stari dvosmerni brojači: A01 | A02 | Sum (Sum = y_A, kao u istoriji); A01 = smer 1x, A02 = smer 2x
        vrednosne = [f"A0{i + 1}" for i in range(len(smerovi_serije))] + ["Sum"] if dvosmerna else ["y_A"]
        zaglavlje = ["Date", "Time"] + vrednosne + ["razlog_A", "oznaka"]
        ws = wb.create_sheet(serija)
        ws.append(zaglavlje)
        if serija in opisi.index:                       # opis lokacije — samo u zaglavlju, drugom bojom
            o = opisi.loc[serija]
            c = ws.cell(row=1, column=len(zaglavlje) + 1,
                        value=f"Lokacija {serija} (mesto {o['mesto']}): {o['naziv']} — {o['smerovi']}")
            c.fill, c.font = _LOKACIJA_FILL, _LOKACIJA_FONT
        for _, r in sr.iterrows():
            oznake = {s: oznake_smera[s].get(r["ts_utc"]) for s in smerovi_serije}
            oznake = {s: o for s, o in oznake.items() if isinstance(o, str)}
            tekst = ", ".join(f"{o} (smer {s})" if dvosmerna else o for s, o in oznake.items())
            y = None if pd.isna(r["y_A"]) else int(r["y_A"])
            if dvosmerna:
                po = [po_smeru[s].get(r["ts_utc"]) for s in smerovi_serije]
                vrednosti = [None if v is None or pd.isna(v) else int(v) for v in po] + [y]
                boje = [oznake.get(s) for s in smerovi_serije] + [next(iter(oznake.values()), None)]
            else:
                vrednosti, boje = [y], [next(iter(oznake.values()), None)]
            ws.append([date.fromisoformat(r["datum"]), _vreme(r["ts"])] + vrednosti + [r["razlog_A"], tekst])
            for i, (v, o) in enumerate(zip(vrednosti, boje)):
                c = ws.cell(row=ws.max_row, column=3 + i)
                if v is None:
                    c.fill = _NEMA_FILL
                elif o:
                    c.fill = _OZNAKA_FILL[o]

        for c in ws[1][:len(zaglavlje)]:
            c.fill, c.font, c.alignment = _ZAGLAVLJE_FILL, _ZAGLAVLJE_FONT, Alignment(horizontal="center")
        for c in ws["A"][1:]:
            c.number_format = "yyyy-mm-dd"
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions[get_column_letter(len(zaglavlje) - 1)].width = 16
        ws.column_dimensions[get_column_letter(len(zaglavlje))].width = 28
        if ws.max_row > 1:
            ws.conditional_formatting.add(
                f"C2:{get_column_letter(2 + len(vrednosne))}{ws.max_row}",
                CellIsRule(operator="greaterThan", formula=[str(PRAG_CRVENO)], fill=_CRVENO_FILL, font=_CRVENO_FONT))
        ws.freeze_panes = "A2"

    info = wb.create_sheet("Info")
    for red in [
        ["Izvor", "UGZINS / Mikrobit QLTCnet Traffic Agent — biciklisti, čitanje svakog sata (XX:02)"],
        ["Time", "kraj sata (01:00 = 00–01 h, 24:00 = 23–24 h)"],
        ["y_A", "broj biciklista u satu (cilj modela, samo bicikli)"],
        ["A01, A02, Sum", "stari dvosmerni brojači: A01 = smer 1x, A02 = smer 2x, Sum = A01 + A02 = y_A"],
        ["razlog_A = ok", "sat je ispravan; inače y_A je prazno, a razlog_A kaže zašto"],
        ["nema_podataka", "uređaj javlja „Nema podataka“ (state 0) — sat nije 0 nego nepoznat"],
        ["nema_citanja", "čitanje u tom satu nije uspelo"],
        ["nema_smera", "u čitanju nema tog smera"],
        ["zastarelo", "podatak brojača je prestar (uređaj ne javlja nove podatke)"],
        ["nepotpun_sat", "razmak čitanja nije ~1 sat"],
        ["negativna_razlika", "brojač je resetovan"],
        ["Crveno", f"vrednost > {PRAG_CRVENO}"],
        ["Sivo", "nema vrednosti za taj sat (razlog u razlog_A)"],
        ["Narandžasto: nula_dan", anomalije.OPIS["nula_dan"]],
        ["Narandžasto: nula_niz", anomalije.OPIS["nula_niz"]],
        ["Žuto: skok", anomalije.OPIS["skok"]],
        ["Napomena", "markirane vrednosti NISU menjane — markiranje samo upozorava; odluka je na korisniku"],
        ["Tamnocrveno", "opis lokacije u zaglavlju sheeta (nije podatak o biciklistima)"],
    ]:
        info.append(red)
    info.append([])
    info.append(["SPISAK BROJAČA"])
    info.cell(row=info.max_row, column=1).font = Font(bold=True)
    info.append(["serija (sheet)", "mesto", "naziv (ulica)", "smerovi", "lat", "lon"])
    for c in info[info.max_row]:
        c.fill, c.font = _LOKACIJA_FILL, _LOKACIJA_FONT
    for _, o in opisi.reset_index().iterrows():
        info.append([o["serija"], o["mesto"], o["naziv"], o["smerovi"], o["lat"], o["lon"]])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 60
    info.column_dimensions["C"].width = 32
    info.column_dimensions["D"].width = 110

    izlaz = folder_excel(koren) / f"bicikli_{d:%Y-%m}.xlsx"
    izlaz.parent.mkdir(parents=True, exist_ok=True)
    wb.save(izlaz)
    log.info("Excel: %s (%d sheetova)", izlaz, len(redosled))
    return izlaz


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--mesec", required=True, help="YYYY-MM")
    a = ap.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    napravi_mesec(date.fromisoformat(a.mesec + "-01"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
