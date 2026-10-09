"""Čuvanje u podaci/YYYY/MM/DD/ i dnevni satni obračun iz tih fajlova."""
import gzip
from datetime import date, datetime, timezone
from pathlib import Path

import pandas as pd

from scraper import citaj, konfig, satno
from tests.test_satno import D, napravi, satni_profil

PRIMER = Path(__file__).parent / "primer_bicikli.xml"


def test_sacuvaj_u_folder_dana(tmp_path, monkeypatch):
    monkeypatch.setattr(konfig, "PODACI", tmp_path)
    data = PRIMER.read_bytes()
    t = datetime(2026, 10, 9, 11, 1, 5, tzinfo=timezone.utc)  # 13:01 lokalno
    csv_put, n, prvo = citaj.sacuvaj(data, t)
    assert csv_put == tmp_path / "2026" / "10" / "09" / "1301.csv" and n == 57 and prvo
    assert gzip.decompress((tmp_path / "2026/10/09/1301.xml.gz").read_bytes()) == data
    assert citaj.postoji_citanje_za_sat(t.replace(minute=20))
    assert not citaj.postoji_citanje_za_sat(t.replace(hour=12))
    _, _, prvo = citaj.sacuvaj(data, t.replace(hour=12))
    assert not prvo
    assert sorted(p.name for p in (tmp_path / "2026/10/09").iterdir()) == \
        ["1301.csv", "1301.xml.gz", "1401.csv", "1401.xml.gz"]


def test_obradi_dan_iz_foldera(tmp_path, monkeypatch):
    monkeypatch.setattr(konfig, "PODACI", tmp_path)
    satni = [satni_profil(24, 10), satni_profil(24, 1)]
    df = napravi(D, satni)
    for vreme, g in df.groupby("vreme_citanja_utc"):
        lok = pd.Timestamp(vreme).tz_convert(konfig.TZ)
        folder = konfig.folder_dana(lok)
        folder.mkdir(parents=True, exist_ok=True)
        g.to_csv(folder / f"{lok:%H%M}.csv", index=False)
    assert satno.dani_sa_citanjima() == [D, date(2026, 10, 6)]
    satno.main(["--dan", D.isoformat()])
    dan = pd.read_csv(tmp_path / "2026/10/05/satno.csv")
    assert len(dan) == 24 and dan["y"].tolist() == [a + b for a, b in zip(*satni)]
    mesec = pd.read_parquet(tmp_path / "2026/10/satno_2026-10.parquet")
    assert len(mesec) == 24
    assert not (tmp_path / "2026/10/05/satno_provera.csv").exists()


def test_excel_mesec(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    test_obradi_dan_iz_foldera(tmp_path, monkeypatch)  # pravi satno.csv za 2026-10-05 i Excel
    wb = load_workbook(tmp_path / "excel" / "bicikli_2026-10.xlsx")
    assert wb.sheetnames == ["26", "Info"]
    ws = wb["26"]
    redovi = list(ws.values)
    assert redovi[0] == ("Date", "Time", "Status", "Sum1", "Sum2", "Sum")
    assert redovi[1][1] == "01:00" and redovi[-1][1] == "24:00" and len(redovi) == 25
    assert redovi[1][3:] == (10, 1, 11) and redovi[1][2] is None
